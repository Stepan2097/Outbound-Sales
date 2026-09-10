#!/usr/bin/env python3
"""Extract DOCX, PDF, and XLSX product sources and load section-level memory."""

from __future__ import annotations

import argparse
import json
import os
import re
import ssl
import urllib.error
import urllib.request
from dataclasses import dataclass
from http.cookiejar import CookieJar
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_CHARS = 6_600


@dataclass
class KnowledgeChunk:
    title: str
    text: str
    source: str
    priority: int


def clean(value: object) -> str:
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_blocks(title: str, blocks: list[str], source: str, priority: int) -> list[KnowledgeChunk]:
    chunks: list[KnowledgeChunk] = []
    current: list[str] = []
    current_size = 0
    part = 1
    for block in [clean(block) for block in blocks if clean(block)]:
        if current and current_size + len(block) + 2 > MAX_CHARS:
            chunk_title = title if part == 1 else f"{title} - part {part}"
            chunks.append(KnowledgeChunk(chunk_title, "\n\n".join(current), source, priority))
            current, current_size, part = [], 0, part + 1
        if len(block) > MAX_CHARS:
            for start in range(0, len(block), MAX_CHARS):
                fragment = block[start : start + MAX_CHARS]
                if current:
                    chunk_title = title if part == 1 else f"{title} - part {part}"
                    chunks.append(KnowledgeChunk(chunk_title, "\n\n".join(current), source, priority))
                    current, current_size, part = [], 0, part + 1
                chunk_title = title if part == 1 else f"{title} - part {part}"
                chunks.append(KnowledgeChunk(chunk_title, fragment, source, priority))
                part += 1
            continue
        current.append(block)
        current_size += len(block) + 2
    if current:
        chunk_title = title if part == 1 else f"{title} - part {part}"
        chunks.append(KnowledgeChunk(chunk_title, "\n\n".join(current), source, priority))
    return chunks


def iter_docx_blocks(document: DocumentType):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        values = [clean(cell.text).replace("\n", " / ") for cell in row.cells]
        if any(values):
            rows.append(" | ".join(values))
    return "\n".join(rows)


def extract_docx(path: Path, priority: int) -> list[KnowledgeChunk]:
    document = Document(path)
    sections: list[tuple[str, list[str]]] = []
    section_title = path.stem
    blocks: list[str] = []
    for block in iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            text = clean(block.text)
            if not text:
                continue
            style = clean(block.style.name).lower()
            if style in {"title", "heading 1"}:
                if blocks:
                    sections.append((section_title, blocks))
                section_title, blocks = text, [text]
            else:
                prefix = "## " if style == "heading 2" else "### " if style == "heading 3" else ""
                blocks.append(f"{prefix}{text}")
        else:
            text = table_text(block)
            if text:
                blocks.append(text)
    if blocks:
        sections.append((section_title, blocks))
    output: list[KnowledgeChunk] = []
    for title, section_blocks in sections:
        output.extend(split_blocks(f"{path.name}: {title}", section_blocks, path.name, priority))
    return output


def extract_pdf(path: Path, priority: int) -> list[KnowledgeChunk]:
    reader = PdfReader(path)
    output: list[KnowledgeChunk] = []
    page_blocks: list[str] = []
    start_page = 1
    size = 0
    for index, page in enumerate(reader.pages, 1):
        text = clean(page.extract_text() or "")
        if page_blocks and size + len(text) > MAX_CHARS:
            output.extend(split_blocks(f"{path.name}: pages {start_page}-{index - 1}", page_blocks, path.name, priority))
            page_blocks, start_page, size = [], index, 0
        if text:
            page_blocks.append(f"Page {index}\n{text}")
            size += len(text)
    if page_blocks:
        output.extend(split_blocks(f"{path.name}: pages {start_page}-{len(reader.pages)}", page_blocks, path.name, priority))
    return output


def extract_xlsx(path: Path, priority: int) -> list[KnowledgeChunk]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    output: list[KnowledgeChunk] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [clean(value).replace("\n", " / ") for value in row]
            while values and not values[-1]:
                values.pop()
            if any(values):
                rows.append(" | ".join(values))
        output.extend(split_blocks(f"{path.name}: {sheet.title}", rows, path.name, priority))
    return output


def knowledge_type(chunk: KnowledgeChunk) -> str:
    text = f"{chunk.title} {chunk.text[:2000]}".lower()
    if any(term in text for term in ["approved proof", "approved claim", "claim discipline", "proof hierarchy"]):
        return "approved_claim"
    if any(term in text for term in ["ideal customer", " icp", "segment hierarchy", "persona map"]):
        return "icp"
    if any(term in text for term in ["objection", "pushback", "competitive landscape", "battlecard"]):
        return "objection"
    if any(term in text for term in ["message", "outreach", "trace framework", "linkedin", "email sequence"]):
        return "winning_outreach"
    if any(term in text for term in ["case study", "customer result"]):
        return "case_study"
    return "product_knowledge"


class ApiClient:
    def __init__(self, base_url: str):
        self.base_url = base_url.rstrip("/")
        jar = CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        self.context = ssl.create_default_context()

    def request(self, path: str, method: str = "GET", payload: dict | None = None):
        data = json.dumps(payload).encode() if payload is not None else None
        request = urllib.request.Request(
            f"{self.base_url}{path}",
            data=data,
            method=method,
            headers={"Content-Type": "application/json", "Accept": "application/json"},
        )
        try:
            with self.opener.open(request, timeout=180) as response:
                return json.loads(response.read().decode() or "{}")
        except urllib.error.HTTPError as error:
            detail = error.read().decode(errors="replace")[:800]
            raise RuntimeError(f"{method} {path} returned HTTP {error.code}: {detail}") from error


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", required=True)
    parser.add_argument("--email", required=True)
    parser.add_argument("--product-id", required=True)
    parser.add_argument("files", nargs="+")
    args = parser.parse_args()
    password = os.environ.get("OUTBOUND_OS_PASSWORD", "")
    if not password:
        raise SystemExit("OUTBOUND_OS_PASSWORD is required.")

    priorities = {
        "AdAction_EMEA_Outbound_Sales_Playbook_2026.docx": 100,
        "AdAction_EMEA_Outbound_Playbook (1).docx": 96,
        "AdAction's Brand Intelligence.pdf": 98,
        "AdAction's Brand Intelligence.docx": 94,
        "AdAction Product Suite.pdf": 97,
        "Solutions Playbook 2026.pdf": 96,
        "Glossary of Terms.xlsx": 92,
    }
    chunks: list[KnowledgeChunk] = []
    for raw_path in args.files:
        path = Path(raw_path)
        priority = priorities.get(path.name, 90)
        if path.suffix.lower() == ".docx":
            chunks.extend(extract_docx(path, priority))
        elif path.suffix.lower() == ".pdf":
            chunks.extend(extract_pdf(path, priority))
        elif path.suffix.lower() in {".xlsx", ".xlsm"}:
            chunks.extend(extract_xlsx(path, priority))
        else:
            raise SystemExit(f"Unsupported file: {path}")

    api = ApiClient(args.base_url)
    api.request("/api/auth/login", "POST", {"email": args.email, "password": password})
    state = api.request("/api/state")
    product = next((item for item in state.get("products", []) if item.get("id") == args.product_id), None)
    if not product:
        raise SystemExit(f"Product not found: {args.product_id}")
    existing_titles = {item.get("title") for item in product.get("knowledge", [])}
    uploaded = 0
    skipped = 0
    for chunk in chunks:
        if chunk.title in existing_titles:
            skipped += 1
            continue
        kind = knowledge_type(chunk)
        api.request(
            "/api/products/knowledge",
            "POST",
            {
                "productId": args.product_id,
                "title": chunk.title,
                "text": chunk.text,
                "type": kind,
                "priority": chunk.priority,
                "tags": f"adaction,source-document,{kind},{Path(chunk.source).stem}",
            },
        )
        existing_titles.add(chunk.title)
        uploaded += 1

    authoritative = [chunk for chunk in chunks if chunk.source == "AdAction_EMEA_Outbound_Sales_Playbook_2026.docx"]
    training_text = clean("\n\n".join(f"# {chunk.title}\n{chunk.text}" for chunk in authoritative))[:29_500]
    if training_text:
        api.request(
            "/api/products/teach",
            "POST",
            {"productId": args.product_id, "forceSelectedProduct": True, "text": training_text},
        )
    final_state = api.request("/api/state")
    final_product = next(item for item in final_state.get("products", []) if item.get("id") == args.product_id)
    print(json.dumps({
        "extracted": len(chunks),
        "uploaded": uploaded,
        "skipped": skipped,
        "storedKnowledge": len(final_product.get("knowledge", [])),
        "memoryStatus": (final_product.get("memory") or {}).get("status"),
        "memoryConfidence": (final_product.get("memory") or {}).get("confidence"),
    }))


if __name__ == "__main__":
    main()
